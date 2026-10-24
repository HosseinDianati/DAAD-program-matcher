"""
Excel Generator matching the exact styling of 'Winter search list- Javad Mahmoodi.xlsx' and 'Winter search Template.xlsx'.
11 Standard Columns:
1. University
2. Program
3. Tuition fee
4. Semester fee
5. Intake
6. Deadline
7. Requirements
8. Application Method
9. Program link
10. Note
11. Apply Link
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Column widths based on client template
COLUMN_WIDTHS = {
    "A": 22.0,  # University
    "B": 24.0,  # Program
    "C": 14.0,  # Tuition fee
    "D": 15.0,  # Semester fee
    "E": 10.0,  # Intake
    "F": 16.0,  # Deadline
    "G": 46.0,  # Requirements
    "H": 15.0,  # Application Method
    "I": 18.0,  # Program link
    "J": 15.0,  # Note
    "K": 20.0,  # Apply Link
}

HEADERS = [
    "University",
    "Program",
    "Tuition fee",
    "Semester fee",
    "Intake",
    "Deadline",
    "Requirements",
    "Application Method",
    "Program link",
    "Note",
    "Apply Link"
]

def create_study_search_workbook(rows: list[dict], title: str = "Germany") -> openpyxl.Workbook:
    """
    Creates an openpyxl Workbook populated with search results.
    rows: list of dicts with keys matching HEADERS.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 1: Title ("Germany")
    title_font = Font(name="Arial", size=12, bold=True)
    title_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")

    ws.cell(row=1, column=1, value=title)
    ws.cell(row=1, column=1).font = title_font
    ws.cell(row=1, column=1).fill = title_fill

    # Fill rest of row 1 with light fill
    for col in range(2, 12):
        cell = ws.cell(row=1, column=col)
        cell.fill = title_fill

    # Row 2: Headers
    header_font = Font(name="Arial", size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    data_font = Font(name="Arial", size=10)
    link_font = Font(name="Arial", size=10, color="0000FF", underline="single")

    for row_idx, row_data in enumerate(rows, start=3):
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row_data.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border

            # Handle text wrap and alignments
            if header == "Requirements":
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.value = val
            elif header in ["Program link", "Apply Link"]:
                cell.alignment = Alignment(horizontal="left", vertical="top")
                if val and str(val).startswith("http"):
                    cell.hyperlink = str(val)
                    cell.font = link_font
                    cell.value = str(val)
                else:
                    cell.value = val
            elif header in ["Tuition fee", "Semester fee", "Intake", "Deadline", "Application Method"]:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                cell.value = val
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                cell.value = val

    # Set column widths
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # Row height formatting
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 28
    for r in range(3, len(rows) + 3):
        ws.row_dimensions[r].height = 75  # Room for multi-line requirements

    return wb


def save_search_to_excel(rows: list[dict], output_path: str, title: str = "Germany"):
    wb = create_study_search_workbook(rows, title)
    wb.save(output_path)
    print(f"Saved {len(rows)} programs to {output_path}")


if __name__ == "__main__":
    # Test sample export
    sample_data = [{
        "University": "Anhalt University of Applied Sciences",
        "Program": "Master of Landscape Architecture (MLA)",
        "Tuition fee": "850 EUR per semester",
        "Semester fee": "96 EUR",
        "Intake": "Winter",
        "Deadline": "15 December - 15 March",
        "Requirements": "1. First degree in landscape architecture, architecture, or urban planning (min. 3 years)\n2. Min 1 year work experience\n3. Portfolio of ~5 projects\n4. English B2 (IELTS 6.5)",
        "Application Method": "Uni-Assist",
        "Program link": "https://www.hs-anhalt.de/mla",
        "Note": "Portfolio required with min 2 professional projects",
        "Apply Link": "https://www.uni-assist.de/en/how-to-apply/apply-online/"
    }]
    import os
    test_out = r"C:\Users\hosse\.gemini\antigravity\brain\0a9935d6-c346-4bea-9876-aed5e1071ea4\scratch\test_output.xlsx"
    save_search_to_excel(sample_data, test_out)
